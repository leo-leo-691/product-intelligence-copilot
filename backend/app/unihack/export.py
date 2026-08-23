"""CSV / XLSX export with exact official delivery headers."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.unihack.schema import require_schema


def write_csv(headers: list[str], rows: list[dict[str, str]]) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.DictWriter(buf, fieldnames=headers, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({h: row.get(h, "") for h in headers})
    return buf.getvalue().encode("utf-8-sig")


def write_xlsx(headers: list[str], rows: list[dict[str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Delivery"
    header_font = Font(name="Calibri", bold=True, color="1C2430")
    fill = PatternFill("solid", fgColor="EDE4CE")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([str(row.get(h, "") or "") for h in headers])
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(28, max(12, len(headers[col - 1]) * 0.7 + 4))
        for cell in ws[letter]:
            cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save_export(path: Path, headers: list[str], rows: list[dict[str, str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        path.write_bytes(write_csv(headers, rows))
    else:
        path.write_bytes(write_xlsx(headers, rows))
    return path


def _csv_header_row(payload: bytes) -> list[str]:
    text = payload.decode("utf-8-sig")
    return next(csv.reader(io.StringIO(text)))


def _xlsx_header_row(payload: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(payload), read_only=True)
    try:
        row = next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
        return [str(c) if c is not None else "" for c in row]
    finally:
        wb.close()


def write_validated_csv(headers: list[str], rows: list[dict[str, str]], expected: list[str]) -> bytes:
    require_schema(headers, expected)
    payload = write_csv(headers, rows)
    require_schema(_csv_header_row(payload), expected)
    return payload


def write_validated_xlsx(headers: list[str], rows: list[dict[str, str]], expected: list[str]) -> bytes:
    require_schema(headers, expected)
    payload = write_xlsx(headers, rows)
    require_schema(_xlsx_header_row(payload), expected)
    return payload
