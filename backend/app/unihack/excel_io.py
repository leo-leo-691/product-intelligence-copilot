"""Read messy Excel/CSV workbooks without assuming row 1 is a clean header."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_HEADERISH = re.compile(r"[A-Za-z]")


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_header_key(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", name.lower())


@dataclass
class SheetTable:
    sheet_name: str
    headers: list[str]
    rows: list[dict[str, str]]
    header_row_index: int
    source_path: str
    duplicates: list[str] = field(default_factory=list)


def _score_header_row(values: list[str]) -> float:
    nonempty = [v for v in values if v]
    if len(nonempty) < 3:
        return 0.0
    unique = len(set(normalize_header_key(v) for v in nonempty if normalize_header_key(v)))
    letterish = sum(1 for v in nonempty if _HEADERISH.search(v))
    numeric = sum(1 for v in nonempty if v.replace(".", "", 1).isdigit())
    if letterish < 3:
        return 0.0
    return unique + letterish * 0.5 - numeric * 0.8


def _detect_header_row(grid: list[list[str]], scan: int = 25) -> int:
    best_i, best_s = 0, -1.0
    for i, row in enumerate(grid[:scan]):
        s = _score_header_row(row)
        if s > best_s:
            best_i, best_s = i, s
    return best_i


def _unique_headers(raw: list[str]) -> tuple[list[str], list[str]]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    duplicates: list[str] = []
    for h in raw:
        name = h or "UNNAMED"
        key = normalize_header_key(name) or name
        n = seen.get(key, 0)
        seen[key] = n + 1
        if n:
            duplicates.append(name)
            headers.append(f"{name}__{n + 1}")
        else:
            headers.append(name)
    return headers, duplicates


def read_csv_table(path: Path) -> SheetTable:
    with path.open(newline="", encoding="utf-8-sig") as f:
        grid = [[cell_str(c) for c in row] for row in csv.reader(f)]
    if not grid:
        return SheetTable(path.stem, [], [], 0, str(path))
    hi = _detect_header_row(grid)
    headers, dups = _unique_headers(grid[hi])
    rows: list[dict[str, str]] = []
    for raw in grid[hi + 1 :]:
        if not any(raw):
            continue
        padded = raw + [""] * (len(headers) - len(raw))
        rows.append({headers[i]: padded[i] for i in range(len(headers))})
    return SheetTable(path.stem, headers, rows, hi, str(path), dups)


def read_xlsx_tables(path: Path, data_only: bool = True) -> list[SheetTable]:
    wb = load_workbook(path, read_only=True, data_only=data_only)
    tables: list[SheetTable] = []
    try:
        for ws in wb.worksheets:
            grid: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                grid.append([cell_str(c) for c in row])
            if not grid:
                continue
            hi = _detect_header_row(grid)
            raw_headers = grid[hi]
            # trim trailing empties
            while raw_headers and not raw_headers[-1]:
                raw_headers = raw_headers[:-1]
            headers, dups = _unique_headers(raw_headers)
            rows: list[dict[str, str]] = []
            for raw in grid[hi + 1 :]:
                if not any(raw):
                    continue
                padded = raw + [""] * (len(headers) - len(raw))
                rows.append({headers[i]: padded[i] for i in range(len(headers))})
            tables.append(
                SheetTable(ws.title, headers, rows, hi, str(path), dups)
            )
    finally:
        wb.close()
    return tables


def read_tables(path: Path) -> list[SheetTable]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [read_csv_table(path)]
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xls"}:
        if suffix == ".xls":
            raise ValueError(
                f"{path.name} is legacy .xls — save as .xlsx or CSV, then re-import."
            )
        return read_xlsx_tables(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def find_column(headers: list[str], wanted: str) -> str | None:
    want = normalize_header_key(wanted)
    for h in headers:
        if normalize_header_key(h) == want:
            return h
    return None
