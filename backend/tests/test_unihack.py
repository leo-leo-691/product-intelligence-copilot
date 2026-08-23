"""UniHack evaluation tests — dynamic rows/columns, exact headers, no invented values."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.unihack.constants import INPUT_ALIASES
from backend.app.unihack.evaluate import evaluate_predictions
from backend.app.unihack.export import write_csv, write_xlsx
from backend.app.unihack.ingest import ingest_input_workbook
from backend.app.unihack.jobs import create_job, run_job
from backend.app.unihack.normalize import clean_input
from backend.app.unihack.official import EXPECTED_OUTPUT_CACHE
from backend.app.unihack.pipeline import process_row
from backend.app.unihack.schema import empty_row, load_delivery_headers, validate_headers

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")


def _xlsx(path: Path, headers: list[str], rows: list[list[object]], sheet: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _csv(path: Path, headers: list[str], rows: list[list[object]]) -> Path:
    import csv

    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    return path


def test_placeholders_are_empty():
    assert clean_input("-- Unbranded --") == ""
    assert clean_input("-- No Unilog Brand --") == ""
    assert clean_input("-- No DIB Brand --") == ""
    assert clean_input("Acme") == "Acme"


def test_ingest_dynamic_row_count(tmp_path: Path):
    headers = list(INPUT_ALIASES)
    rows = [[f"SKU-{i}", f"desc {i}", "Acme", "", "", "Acme"] for i in range(7)]
    data = ingest_input_workbook(_xlsx(tmp_path / "in.xlsx", headers, rows))
    assert data["row_count"] == 7


def test_ingest_dynamic_and_unknown_columns(tmp_path: Path):
    headers = ["Mfg_Part_Num", "Part_Desc", "Custom_Col", "E1_Brand"]
    data = ingest_input_workbook(
        _xlsx(tmp_path / "in.xlsx", headers, [["A1", "valve", "keep-me", "-- Unbranded --"]])
    )
    assert data["row_count"] == 1
    row = data["rows"][0]
    assert row["Custom_Col"] == "keep-me"
    assert row["E1_Brand"] == ""
    assert "Unilog_Brand" not in headers
    assert row.get("Unilog_Brand", "") == ""


def test_ingest_missing_input_values(tmp_path: Path):
    headers = ["Mfg_Part_Num", "Part_Desc"]
    data = ingest_input_workbook(_xlsx(tmp_path / "in.xlsx", headers, [["SKU", ""]]))
    assert data["rows"][0]["Part_Desc"] == ""
    assert data["rows"][0]["Mfg_Part_Num"] == "SKU"


def test_official_expected_output_headers_loaded():
    if not EXPECTED_OUTPUT_CACHE.exists():
        pytest.skip("official expected output CSV not cached")
    loaded = load_delivery_headers(EXPECTED_OUTPUT_CACHE)
    assert loaded["header_count"] > 0
    assert loaded["headers"][0] == "MFR URL"
    assert "Mfg_Part_Num" in loaded["headers"]
    check = validate_headers(loaded["headers"], loaded["headers"])
    assert check["valid"] is True


def test_process_row_emits_every_header_and_does_not_hallucinate(tmp_path: Path):
    headers = ["MFR URL", "Mfg_Part_Num", "Part_Desc", "Mystery_Spec", "Part_Manuf"]
    raw = {
        "row_number": 1,
        "Mfg_Part_Num": "ABC-9",
        "Part_Desc": "2 in steel elbow",
        "Part_Manuf": "UnknownMfrXYZ",
    }
    out = process_row(raw, headers, manufacturer_path=tmp_path / "missing.xlsx", enrichment_enabled=False)
    assert list(out["delivery"].keys()) == headers
    assert out["delivery"]["Mfg_Part_Num"] == "ABC-9"
    assert out["delivery"]["Mystery_Spec"] == ""
    assert out["delivery"]["MFR URL"] == ""
    assert "titanium" not in str(out["delivery"]).lower()


def test_evaluation_exact_match_and_missing():
    headers = ["Mfg_Part_Num", "Color"]
    predicted = [{"Mfg_Part_Num": "A1", "Color": "Red"}]
    expected = [{"mfg_part_num": "A1", "expected": {"Mfg_Part_Num": "A1", "Color": "Red"}}]
    metrics = evaluate_predictions(predicted, expected, headers)
    assert metrics["fields_correct"] == 2
    assert metrics["field_level_accuracy_pct"] == 100.0
    predicted2 = [{"Mfg_Part_Num": "A1", "Color": ""}]
    metrics2 = evaluate_predictions(predicted2, expected, headers)
    assert metrics2["fields_missing"] >= 1


def test_csv_and_xlsx_preserve_exact_headers():
    headers = ["MFR URL", "Odd Header / 2", "Mfg_Part_Num"]
    rows = [empty_row(headers)]
    rows[0]["Mfg_Part_Num"] = "Z1"
    csv_text = write_csv(headers, rows).decode("utf-8-sig")
    assert csv_text.splitlines()[0] == "MFR URL,Odd Header / 2,Mfg_Part_Num"
    xlsx = write_xlsx(headers, rows)
    dest = Path("_tmp_unihack_hdr.xlsx")
    try:
        dest.write_bytes(xlsx)
        wb = load_workbook(dest)
        got = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
        assert got == headers
    finally:
        dest.unlink(missing_ok=True)


def test_job_row_failure_isolation(tmp_path: Path, monkeypatch):
    headers = ["Mfg_Part_Num", "Part_Desc"]
    schema = _csv(tmp_path / "schema.csv", headers, [])
    inp = _xlsx(
        tmp_path / "in.xlsx",
        ["Mfg_Part_Num", "Part_Desc"],
        [["OK-1", "valve"], ["OK-2", "sensor"]],
    )
    from backend.app.unihack import jobs as jobs_mod
    from backend.app.unihack.pipeline import process_row as real

    def flaky(raw, hdrs, **kwargs):
        if raw.get("Mfg_Part_Num") == "OK-1":
            raise RuntimeError("boom")
        return real(raw, hdrs, enrichment_enabled=False, **kwargs)

    monkeypatch.setattr(jobs_mod, "process_row", flaky)
    monkeypatch.setattr(jobs_mod, "find_file", lambda kind: inp if kind == "input_1000" else None)
    monkeypatch.setattr(jobs_mod, "find_delivery_schema", lambda: schema)
    monkeypatch.setattr(jobs_mod, "ensure_expected_output", lambda: schema)
    monkeypatch.setattr(jobs_mod, "ensure_sample_input", lambda: inp)
    job = create_job(input_path=inp, evaluate=False)
    finished = run_job(job["id"])
    assert finished["status"] == "PARTIAL"
    assert finished["progress"]["successful"] == 1
    assert finished["progress"]["failed"] == 1
    from backend.app.unihack.store import list_rows

    rows = list_rows(finished["id"])
    failed = next(r for r in rows if r.get("status") == "error")
    assert list(failed["delivery"].keys()) == headers


def test_output_generation_blank_unknown(tmp_path: Path):
    headers = ["Mfg_Part_Num", "UPC", "Country Of Origin"]
    raw = {"row_number": 1, "Mfg_Part_Num": "N1", "Part_Desc": "nut"}
    out = process_row(raw, headers, enrichment_enabled=False)
    assert out["delivery"]["UPC"] == ""
    assert out["delivery"]["Country Of Origin"] == ""
    assert out["delivery"]["Mfg_Part_Num"] == "N1"


# --- Enrichment architecture tests (no hardcoded sample products) ---

_SAMPLE_HTML = """
<html><head>
<script type="application/ld+json">
{"@type":"Product","name":"Pro Washer","brand":{"name":"AcmePro"},"mpn":"ZZ-999-XY","description":"24 inch unit"}
</script></head><body>
<dl><dt>Sound Level</dt><dd>39 dBA</dd><dt>Voltage</dt><dd>120 V</dd></dl>
<ul class="features"><li>Third Rack</li><li>Stainless Steel Tub</li></ul>
<img src="https://cdn.example.com/zz999.jpg" alt="product photo"/>
<a href="/specs/ZZ-999-XY.pdf">Specification Sheet</a>
</body></html>
"""


def _mock_search_fn(query, max_results=5):
    return [
        (
            "https://manufacturer.example.com/p/ZZ-999-XY",
            "AcmePro ZZ-999-XY Washer",
            "ZZ-999-XY 24 inch washer 39 dBA",
        )
    ]


class _MockRetriever:
    def get(self, url, **kwargs):
        from backend.app.unihack.retrieve import RetrievedPage

        return RetrievedPage(
            url=url,
            html=_SAMPLE_HTML.replace("ZZ-999-XY", kwargs.get("part_number") or "ZZ-999-XY"),
            text="ZZ-999-XY 39 dBA 120 V Third Rack Stainless Steel Tub",
            content_type="text/html",
            status_code=200,
        )


def test_identify_product_from_arbitrary_columns():
    from backend.app.unihack.identify import identify_product, normalize_mpn

    identity = identify_product(
        {
            "Mfg_Part_Num": "  ab-12/x ",
            "Part_Desc": "2 in steel elbow",
            "Part_Manuf": "Acme Corp",
            "E1_Brand": "-- Unbranded --",
        }
    )
    assert identity.mfg_part_num == "ab-12/x"
    assert normalize_mpn("ab-12/x") == "AB-12/X"
    assert identity.manufacturer_query == "Acme Corp"
    assert identity.brand_query == ""
    assert any("ab-12" in q.lower() for q in identity.search_queries)


def test_enrichment_extracts_from_mock_source(tmp_path: Path):
    headers = [
        "MFR URL",
        "Mfg_Part_Num",
        "MANUFACTURER_NAME",
        "BRAND_NAME",
        "ATTRIBUTE_LABEL 1",
        "ATTRIBUTE_VALUE 1",
        "ATTRIBUTE_UOM 1",
        "ITEM_FEATURES_1",
        "Product Image",
        "Actual Image (Yes/No)",
        "Specification Sheet",
    ]
    raw = {
        "row_number": 1,
        "Mfg_Part_Num": "ZZ-999-XY",
        "Part_Desc": "Washer SS",
        "Part_Manuf": "Acme Corp",
    }
    from backend.app.unihack.enrich import enrich_product
    from backend.app.unihack.sources import discover_sources

    result = enrich_product(
        raw,
        headers,
        retriever=_MockRetriever(),
        search_fn=_mock_search_fn,
    )
    assert result.attributes_extracted >= 3
    assert result.delivery.get("BRAND_NAME") == "AcmePro"
    assert "39 dBA" in str(result.delivery.values())
    assert result.delivery.get("Actual Image (Yes/No)") == "Yes"
    assert result.delivery.get("Mystery_Spec", "") == ""


def test_enrichment_does_not_hallucinate_unknown_specs():
    from backend.app.unihack.enrich import enrich_product

    headers = ["Mfg_Part_Num", "Voltage", "ATTRIBUTE_VALUE 1"]
    raw = {"row_number": 2, "Mfg_Part_Num": "EMPTY-404", "Part_Desc": "unknown gadget"}

    def no_sources(query, max_results=5):
        return []

    result = enrich_product(raw, headers, search_fn=no_sources)
    assert result.attributes_extracted == 0
    assert result.delivery.get("Voltage", "") == ""
    assert "120" not in str(result.delivery)


def test_no_hardcoded_sample_product_dictionary():
    """Ensure PDSH4816AF / WDTS7024RZ are not baked into source code."""
    from backend.app.unihack import enrich, extract_attrs, pipeline, sources

    for mod in (enrich, extract_attrs, pipeline, sources):
        text = Path(mod.__file__).read_text(encoding="utf-8")
        assert "PDSH4816AF" not in text
        assert "WDTS7024RZ" not in text


@pytest.mark.parametrize(
    "raw",
    [
        {"Mfg_Part_Num": "FAKE-VALVE-001", "Part_Desc": "ball valve", "Part_Manuf": "UnknownMfrXYZ"},
        {"Mfg_Part_Num": "???", "Part_Desc": ""},
        {"Part_Desc": "generic part"},
        {"Mfg_Part_Num": "DUP-1", "Part_Desc": "item", "Extra_Col": "ignored"},
    ],
)
def test_arbitrary_products_row_isolation(raw):
    headers = ["Mfg_Part_Num", "Part_Desc", "UPC", "Voltage"]
    raw = {"row_number": 99, **raw}
    out = process_row(raw, headers, enrichment_enabled=False)
    assert list(out["delivery"].keys()) == headers
    assert out["status"] == "ok"
    assert out["delivery"].get("Voltage", "") == ""


def test_page_cache_avoids_duplicate_keys(tmp_path: Path, monkeypatch):
    from backend.app.unihack import cache as cache_mod

    monkeypatch.setattr(cache_mod, "CACHE_DIR", tmp_path)
    cache_mod.set_cached("Acme", "P1", "https://example.com/p1", {"html": "<html/>"})
    hit = cache_mod.get_cached("Acme", "P1", "https://example.com/p1")
    assert hit is not None
    assert hit["html"] == "<html/>"
    miss = cache_mod.get_cached("Acme", "P2", "https://example.com/p1")
    assert miss is None


def test_verification_products_with_mocked_manufacturer_pages():
    """PDSH4816AF / WDTS7024RZ used only as verification — not hardcoded."""
    cases = {
        "PDSH4816AF": (
            "Frigidaire",
            '<html><body><script type="application/ld+json">'
            '{"@type":"Product","brand":{"name":"Frigidaire Professional"},"mpn":"PDSH4816AF",'
            '"description":"24 inch dishwasher"}</script>'
            "<dl><dt>Sound Level</dt><dd>49 dBA</dd><dt>Voltage</dt><dd>120 V</dd></dl>"
            '<li>Third Rack</li><img src="https://cdn.example/frig.jpg"/></body></html>',
        ),
        "WDTS7024RZ": (
            "Whirlpool",
            '<html><body><script type="application/ld+json">'
            '{"@type":"Product","brand":{"name":"Whirlpool"},"mpn":"WDTS7024RZ",'
            '"description":"24 inch dishwasher"}</script>'
            "<dl><dt>Sound Level</dt><dd>47 dBA</dd></dl>"
            '<li>Quiet Operation</li></body></html>',
        ),
    }
    headers = [
        "MFR URL",
        "Mfg_Part_Num",
        "MANUFACTURER_NAME",
        "BRAND_NAME",
        "ATTRIBUTE_LABEL 1",
        "ATTRIBUTE_VALUE 1",
        "ITEM_FEATURES_1",
        "Product Image",
    ]
    from backend.app.unihack.enrich import enrich_product

    counts: dict[str, int] = {}
    for mpn, (mfr, html) in cases.items():

        def make_search(part=mpn, manufacturer=mfr):
            def _search(query, max_results=5):
                return [
                    (
                        f"https://{manufacturer.lower()}.example.com/{part}",
                        f"{manufacturer} {part}",
                        part,
                    )
                ]

            return _search

        class Retriever:
            def get(self, url, **kwargs):
                from backend.app.unihack.retrieve import RetrievedPage

                return RetrievedPage(
                    url=url,
                    html=html,
                    text=html,
                    content_type="text/html",
                    status_code=200,
                )

        raw = {
            "row_number": 1,
            "Mfg_Part_Num": mpn,
            "Part_Desc": f"{mpn} Dishwasher SS - Display Only",
            "Part_Manuf": "Appliance Dealers Cooperative (APPDE)",
        }
        result = enrich_product(raw, headers, retriever=Retriever(), search_fn=make_search())
        counts[mpn] = result.attributes_extracted
        assert result.attributes_extracted >= 4
        assert result.delivery.get("BRAND_NAME")

    assert counts["PDSH4816AF"] >= 4
    assert counts["WDTS7024RZ"] >= 3


# ============================================================
# New enrichment architecture tests
# ============================================================


def test_jsonld_nested_graph_extraction():
    """JSON-LD with @graph wrapper and nested structures."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><head><script type="application/ld+json">
    {"@graph": [
        {"@type": "Product", "name": "Super Faucet", "brand": {"name": "DeltaFlow"},
         "mpn": "DF-100", "sku": "SKU-DF100",
         "manufacturer": {"name": "Delta Corp", "url": "https://delta.example.com"},
         "gtin13": "1234567890123",
         "image": ["https://cdn.example/img1.jpg", "https://cdn.example/img2.jpg"],
         "offers": {"price": "299.99"},
         "additionalProperty": [
             {"name": "Flow Rate", "value": "1.8 GPM"},
             {"name": "Connection Size", "value": "3/8 inch"}
         ]}
    ]}</script></head><body><p>DF-100 faucet specs</p></body></html>"""

    identity = ProductIdentity(
        mfg_part_num="DF-100",
        part_desc="faucet",
        manufacturer_query="Delta Corp",
        brand_query="DeltaFlow",
        mpn_normalized="DF100",
    )
    page = RetrievedPage(url="https://delta.example.com/p/DF-100", html=html, text="DF-100 faucet",
                         content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    keys = {a.key for a in result.attributes}
    assert "brand" in keys
    assert "manufacturer" in keys
    assert "sku" in keys
    assert "gtin" in keys
    assert "product_name" in keys
    assert len(result.attributes) >= 5
    assert len(result.images) >= 2


def test_microdata_extraction():
    """Schema.org microdata via itemscope/itemprop."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><body>
    <div itemscope itemtype="https://schema.org/Product">
        <span itemprop="name">Brass Elbow Fitting</span>
        <span itemprop="brand">FitPro</span>
        <span itemprop="mpn">FP-90-BRS</span>
        <span itemprop="sku">SKU-FP90</span>
        <meta itemprop="gtin13" content="9876543210123"/>
        <img itemprop="image" src="https://cdn.example/fp90.jpg"/>
        <span itemprop="description">90 degree brass elbow</span>
    </div>
    </body></html>"""

    identity = ProductIdentity(
        mfg_part_num="FP-90-BRS",
        part_desc="brass elbow",
        manufacturer_query="FitPro",
        brand_query="",
        mpn_normalized="FP90BRS",
    )
    page = RetrievedPage(url="https://fitpro.example.com/p/FP-90-BRS", html=html,
                         text="FP-90-BRS brass elbow", content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    keys = {a.key for a in result.attributes}
    assert "brand" in keys
    assert "part_number" in keys
    assert "sku" in keys
    assert len(result.images) >= 1


def test_meta_tag_extraction():
    """OpenGraph and meta tag extraction."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><head>
    <meta property="og:title" content="Premium Water Heater WH-50"/>
    <meta property="og:description" content="50 gallon natural gas water heater"/>
    <meta property="og:image" content="https://cdn.example/wh50.jpg"/>
    <meta property="og:site_name" content="RheemPro"/>
    <meta name="description" content="High efficiency 50 gallon water heater with Energy Star"/>
    </head><body><p>WH-50 water heater details</p></body></html>"""

    identity = ProductIdentity(
        mfg_part_num="WH-50",
        part_desc="water heater",
        manufacturer_query="RheemPro",
        brand_query="",
        mpn_normalized="WH50",
    )
    page = RetrievedPage(url="https://rheem.example.com/wh50", html=html,
                         text="WH-50 water heater details", content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    keys = {a.key for a in result.attributes}
    assert "product_name" in keys
    assert "manufacturer" in keys
    assert len(result.images) >= 1


def test_spec_table_extraction():
    """Dedicated specification table parsing."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><body>
    <div class="product-specifications">
        <table>
            <tr><th>Voltage</th><td>240 V</td></tr>
            <tr><th>Amperage</th><td>30 A</td></tr>
            <tr><th>Width</th><td>24 in</td></tr>
            <tr><th>Height</th><td>34 in</td></tr>
            <tr><th>Weight</th><td>150 lbs</td></tr>
            <tr><th>Country of Origin</th><td>USA</td></tr>
            <tr><th>UPC</th><td>012345678901</td></tr>
            <tr><th>Warranty</th><td>5 year limited warranty</td></tr>
            <tr><th>Certifications</th><td>UL Listed, NSF</td></tr>
            <tr><th>Color/Finish</th><td>Stainless Steel</td></tr>
            <tr><th>Application</th><td>Residential</td></tr>
        </table>
    </div>
    <p>TEST-SPEC-001 is a great product</p>
    </body></html>"""

    identity = ProductIdentity(
        mfg_part_num="TEST-SPEC-001",
        part_desc="appliance",
        manufacturer_query="TestMfr",
        brand_query="",
        mpn_normalized="TESTSPEC001",
    )
    page = RetrievedPage(url="https://testmfr.example.com/p/TEST-SPEC-001", html=html,
                         text="TEST-SPEC-001 appliance 240 V 30 A 24 in 34 in 150 lbs USA",
                         content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    keys = {a.key for a in result.attributes}
    assert "voltage" in keys
    assert "amperage" in keys
    assert "country_of_origin" in keys
    assert "upc" in keys
    assert "warranty" in keys
    assert "certification" in keys
    assert "color" in keys
    assert "application" in keys
    assert len(result.attributes) >= 8


def test_document_link_extraction():
    """Document link detection for PDFs and guides."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><body>
    <p>DOC-TEST-001 product page</p>
    <a href="/docs/spec-sheet.pdf">Specification Sheet PDF</a>
    <a href="/docs/install-guide.pdf">Installation Manual</a>
    <a href="/docs/owners-manual.pdf">Owner's Manual</a>
    <a href="/docs/sds-data.pdf">Safety Data Sheet</a>
    <a href="/docs/warranty-info.pdf">Warranty Information</a>
    <a href="/docs/submittal.pdf">Submittal Sheet</a>
    </body></html>"""

    identity = ProductIdentity(
        mfg_part_num="DOC-TEST-001",
        part_desc="test product",
        manufacturer_query="TestCo",
        brand_query="",
        mpn_normalized="DOCTEST001",
    )
    page = RetrievedPage(url="https://testco.example.com/p/DOC-TEST-001", html=html,
                         text="DOC-TEST-001 product page", content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    assert result.spec_sheet_url != ""
    assert len(result.documents) >= 3
    assert "installation_manual" in result.documents
    assert "sds" in result.documents


def test_feature_extraction_from_containers():
    """Feature extraction from product-features containers."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><body>
    <p>FEAT-TEST-001 product</p>
    <div class="product-features">
        <ul>
            <li>Energy Star Certified</li>
            <li>Fingerprint Resistant Stainless Steel</li>
            <li>Third Level Rack for Flatware</li>
            <li>EvenDry System</li>
            <li>OrbitClean Wash System</li>
        </ul>
    </div>
    <div id="features">
        <ul>
            <li>Dishwasher Safe Racks</li>
        </ul>
    </div>
    </body></html>"""

    identity = ProductIdentity(
        mfg_part_num="FEAT-TEST-001",
        part_desc="dishwasher",
        manufacturer_query="TestBrand",
        brand_query="",
        mpn_normalized="FEATTEST001",
    )
    page = RetrievedPage(url="https://test.example.com/p/FEAT-TEST-001", html=html,
                         text="FEAT-TEST-001 dishwasher product", content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    feature_values = [a.value for a in result.attributes if a.key == "feature"]
    assert len(feature_values) >= 4
    assert any("Energy Star" in f for f in feature_values)


def test_image_extraction_gallery_priority():
    """Image extraction with gallery container priority."""
    from backend.app.unihack.extract_attrs import extract_from_page
    from backend.app.unihack.identify import ProductIdentity
    from backend.app.unihack.retrieve import RetrievedPage

    html = """<html><body>
    <p>IMG-TEST-001 product</p>
    <img src="https://cdn.example/logo.png" alt="company logo"/>
    <div class="product-gallery">
        <img src="https://cdn.example/product1.jpg" alt="product front view"/>
        <img src="https://cdn.example/product2.jpg" alt="product side view"/>
    </div>
    <img src="https://cdn.example/other.jpg" alt="lifestyle photo"/>
    <img src="data:image/gif;base64,R0lGODlhAQABAIAAAP" alt="pixel"/>
    </body></html>"""

    identity = ProductIdentity(
        mfg_part_num="IMG-TEST-001",
        part_desc="product",
        manufacturer_query="TestCo",
        brand_query="",
        mpn_normalized="IMGTEST001",
    )
    page = RetrievedPage(url="https://test.example.com/p/IMG-TEST-001", html=html,
                         text="IMG-TEST-001 product", content_type="text/html", status_code=200)
    result = extract_from_page(page, identity)

    # Should include gallery images, skip logo and data: URIs
    assert len(result.images) >= 2
    assert all("data:" not in img for img in result.images)
    assert all("logo" not in img for img in result.images)


def test_252_column_mapping_with_documents():
    """Verify document URLs map to correct delivery headers."""
    from backend.app.unihack.attribute_map import map_to_delivery
    from backend.app.unihack.extract_attrs import ExtractionResult, ExtractedAttribute
    from backend.app.unihack.identify import ProductIdentity

    headers = [
        "Mfg_Part_Num",
        "BRAND_NAME",
        "MANUFACTURER_NAME",
        "Specification Sheet",
        "SDS",
        "Instruction/Installation Manual",
        "Owners/User Manual",
        "Warranty Information",
        "Country Of Origin",
        "UPC",
        "Product Image",
        "Alternate Image 1",
        "Alternate Image 2",
        "Actual Image (Yes/No)",
        "Application",
        "Standard/Approvals",
        "TRADE_NAME",
        "ALTERNATE_PART_NUMBER",
    ]

    extraction = ExtractionResult(
        attributes=[
            ExtractedAttribute(key="brand", value="AcmeBrand", source_url="http://test.com", evidence="LD"),
            ExtractedAttribute(key="manufacturer", value="Acme Corp", source_url="http://test.com", evidence="LD"),
            ExtractedAttribute(key="country_of_origin", value="USA", source_url="http://test.com", evidence="table"),
            ExtractedAttribute(key="upc", value="012345678901", source_url="http://test.com", evidence="LD"),
            ExtractedAttribute(key="application", value="Commercial", source_url="http://test.com", evidence="table"),
            ExtractedAttribute(key="certification", value="UL, NSF", source_url="http://test.com", evidence="table"),
            ExtractedAttribute(key="trade_name", value="AcmePro Series", source_url="http://test.com", evidence="LD"),
            ExtractedAttribute(key="alternate_part_number", value="ACME-ALT-1", source_url="http://test.com", evidence="table"),
        ],
        images=["https://cdn.example/img1.jpg", "https://cdn.example/img2.jpg", "https://cdn.example/img3.jpg"],
        spec_sheet_url="https://example.com/spec.pdf",
        documents={
            "sds": "https://example.com/sds.pdf",
            "installation_manual": "https://example.com/install.pdf",
            "owners_manual": "https://example.com/owners.pdf",
            "warranty_doc": "https://example.com/warranty.pdf",
        },
    )

    identity = ProductIdentity(
        mfg_part_num="TEST-MAP-001",
        part_desc="test product",
        manufacturer_query="Acme Corp",
        brand_query="AcmeBrand",
        mpn_normalized="TESTMAP001",
    )

    delivery, provenance = map_to_delivery(extraction, headers, identity)

    assert delivery.get("BRAND_NAME") == "AcmeBrand"
    assert delivery.get("MANUFACTURER_NAME") == "Acme Corp"
    assert delivery.get("Country Of Origin") == "USA"
    assert delivery.get("UPC") == "012345678901"
    assert delivery.get("Specification Sheet") == "https://example.com/spec.pdf"
    assert delivery.get("SDS") == "https://example.com/sds.pdf"
    assert delivery.get("Instruction/Installation Manual") == "https://example.com/install.pdf"
    assert delivery.get("Actual Image (Yes/No)") == "Yes"
    assert delivery.get("Product Image") == "https://cdn.example/img1.jpg"
    assert delivery.get("Application") == "Commercial"
    assert delivery.get("Standard/Approvals") == "UL, NSF"
    assert delivery.get("TRADE_NAME") == "AcmePro Series"
    assert delivery.get("ALTERNATE_PART_NUMBER") == "ACME-ALT-1"


def test_search_cache():
    """Search result caching avoids repeated queries."""
    from backend.app.unihack.cache import get_search_cached, set_search_cached

    results = [("https://example.com/p1", "Product 1", "snippet 1")]
    set_search_cached("test query 123", results)
    cached = get_search_cached("test query 123")
    assert cached is not None
    assert cached[0][0] == "https://example.com/p1"
    # Different query should miss
    miss = get_search_cached("different query")
    assert miss is None


def test_label_normalization():
    """Verify label normalization maps various phrasings correctly."""
    from backend.app.unihack.extract_attrs import _label_to_key

    assert _label_to_key("Sound Level") == "sound_level"
    assert _label_to_key("Sound Level (dBA)") == "sound_level"
    assert _label_to_key("Noise Level") == "sound_level"
    assert _label_to_key("Operating Voltage") == "voltage"
    assert _label_to_key("Voltage Rating") == "voltage"
    assert _label_to_key("Country of Origin") == "country_of_origin"
    assert _label_to_key("Country of Manufacture") == "country_of_origin"
    assert _label_to_key("UPC Code") == "upc"
    assert _label_to_key("Manufacturer Part Number") == "part_number"
    assert _label_to_key("Certifications") == "certification"
    assert _label_to_key("Energy Star Certified") == "energy_star"


@pytest.mark.parametrize(
    "raw,category",
    [
        ({"Mfg_Part_Num": "PLUMB-001", "Part_Desc": "brass ball valve 1 inch", "Part_Manuf": "Watts"}, "plumbing"),
        ({"Mfg_Part_Num": "HVAC-001", "Part_Desc": "furnace filter 20x25", "Part_Manuf": "Lennox"}, "hvac"),
        ({"Mfg_Part_Num": "ELEC-001", "Part_Desc": "circuit breaker 20A", "Part_Manuf": "Eaton"}, "electrical"),
        ({"Mfg_Part_Num": "TOOL-001", "Part_Desc": "cordless drill 20V", "Part_Manuf": "DeWalt"}, "tools"),
    ],
)
def test_arbitrary_category_products(raw, category):
    """Test products from different categories — no category-specific hardcoding."""
    headers = ["Mfg_Part_Num", "Part_Desc", "Part_Manuf", "UPC", "Voltage", "BRAND_NAME"]
    raw = {"row_number": 42, **raw}
    out = process_row(raw, headers, enrichment_enabled=False)
    assert list(out["delivery"].keys()) == headers
    assert out["status"] == "ok"
    # No hallucinated values for unknown fields
    assert out["delivery"].get("UPC", "") == ""
    assert out["delivery"].get("Voltage", "") == ""


def test_fake_product_no_hallucination():
    """A completely fake product should produce no invented specifications."""
    headers = [
        "Mfg_Part_Num", "Part_Desc", "BRAND_NAME", "UPC", "Voltage",
        "ATTRIBUTE_VALUE 1", "Product Image", "Specification Sheet",
        "Country Of Origin", "Standard/Approvals",
    ]
    raw = {
        "row_number": 1,
        "Mfg_Part_Num": "FAKE-VALVE-001",
        "Part_Desc": "nonexistent product xyz",
        "Part_Manuf": "FakeManufacturerCorp",
    }

    def no_sources(query, max_results=5):
        return []

    from backend.app.unihack.enrich import enrich_product

    result = enrich_product(raw, headers, search_fn=no_sources)
    assert result.attributes_extracted == 0
    # All spec fields should be empty
    for key in ["Voltage", "ATTRIBUTE_VALUE 1", "Product Image",
                "Specification Sheet", "Country Of Origin", "Standard/Approvals"]:
        assert result.delivery.get(key, "") == ""


def test_expanded_verification_with_rich_mock():
    """Verification with rich mock data to test comprehensive extraction."""
    from backend.app.unihack.enrich import enrich_product

    rich_html = """<html><head>
    <meta property="og:title" content="Professional Dishwasher RICH-001"/>
    <meta property="og:image" content="https://cdn.example/og-image.jpg"/>
    <meta property="og:site_name" content="TestBrand"/>
    <script type="application/ld+json">
    {"@type":"Product","name":"Professional Dishwasher RICH-001",
     "brand":{"name":"TestBrand Pro"},
     "manufacturer":{"name":"TestManufacturer","url":"https://testmfr.com"},
     "mpn":"RICH-001","sku":"SKU-RICH001",
     "gtin13":"1234567890123",
     "image":["https://cdn.example/main.jpg","https://cdn.example/alt1.jpg"],
     "description":"Professional grade 24 inch built-in dishwasher",
     "offers":{"price":"899.99"}}
    </script>
    </head><body>
    <div class="product-specifications">
        <table>
            <tr><th>Sound Level</th><td>41 dBA</td></tr>
            <tr><th>Voltage</th><td>120 V</td></tr>
            <tr><th>Amperage</th><td>15 A</td></tr>
            <tr><th>Width</th><td>24 inches</td></tr>
            <tr><th>Height</th><td>34 inches</td></tr>
            <tr><th>Depth</th><td>24 inches</td></tr>
            <tr><th>Weight</th><td>80 lbs</td></tr>
            <tr><th>Wash Cycles</th><td>5</td></tr>
            <tr><th>Country of Origin</th><td>USA</td></tr>
            <tr><th>Color/Finish</th><td>Fingerprint Resistant Stainless Steel</td></tr>
            <tr><th>Certifications</th><td>Energy Star, NSF, UL</td></tr>
            <tr><th>Warranty</th><td>2 year limited warranty</td></tr>
            <tr><th>Application</th><td>Residential</td></tr>
        </table>
    </div>
    <div class="product-features">
        <ul>
            <li>Third Level Rack for Flatware</li>
            <li>EvenDry System</li>
            <li>OrbitClean Wash System</li>
            <li>Fingerprint Resistant Stainless Steel</li>
            <li>Energy Star Certified</li>
        </ul>
    </div>
    <a href="/docs/RICH-001-spec.pdf">Specification Sheet</a>
    <a href="/docs/RICH-001-install.pdf">Installation Manual</a>
    <a href="/docs/RICH-001-sds.pdf">Safety Data Sheet</a>
    </body></html>"""

    headers = [
        "MFR URL", "Mfg_Part_Num", "MANUFACTURER_NAME", "BRAND_NAME",
        "SKU - MY_PART_NUMBER", "GTIN", "Product Name",
        "ATTRIBUTE_LABEL 1", "ATTRIBUTE_VALUE 1", "ATTRIBUTE_UOM 1",
        "ATTRIBUTE_LABEL 2", "ATTRIBUTE_VALUE 2", "ATTRIBUTE_UOM 2",
        "ATTRIBUTE_LABEL 3", "ATTRIBUTE_VALUE 3", "ATTRIBUTE_UOM 3",
        "WIDTH", "WIDTH_UOM", "HEIGHT", "HEIGHT_UOM", "LENGTH", "LENGTH_UOM",
        "WEIGHT", "WEIGHT_UOM",
        "ITEM_FEATURES_1", "ITEM_FEATURES_2", "ITEM_FEATURES_3",
        "Product Image", "Alternate Image 1",
        "Actual Image (Yes/No)",
        "Specification Sheet", "SDS",
        "Instruction/Installation Manual",
        "Country Of Origin", "Standard/Approvals",
        "Warranty", "Application", "With",
        "UPC", "List Price",
    ]

    def search_fn(query, max_results=5):
        return [("https://testmfr.example.com/RICH-001", "TestBrand Pro RICH-001", "RICH-001")]

    class Retriever:
        def get(self, url, **kwargs):
            from backend.app.unihack.retrieve import RetrievedPage
            return RetrievedPage(
                url=url, html=rich_html, text=rich_html,
                content_type="text/html", status_code=200,
            )

    raw = {
        "row_number": 1,
        "Mfg_Part_Num": "RICH-001",
        "Part_Desc": "Professional Dishwasher",
        "Part_Manuf": "TestManufacturer",
    }

    result = enrich_product(raw, headers, retriever=Retriever(), search_fn=search_fn)

    # Should extract MANY attributes
    assert result.attributes_extracted >= 10, f"Only extracted {result.attributes_extracted} attributes"
    assert result.images_found >= 2
    assert result.spec_sheet_found is True

    # Verify key deliverables
    d = result.delivery
    assert d.get("BRAND_NAME") == "TestBrand Pro"
    assert d.get("MANUFACTURER_NAME") == "TestManufacturer"
    assert d.get("Country Of Origin") == "USA"
    assert d.get("Standard/Approvals") == "Energy Star, NSF, UL"
    assert d.get("Actual Image (Yes/No)") == "Yes"
    assert d.get("Specification Sheet") != ""
    assert d.get("Product Image") != ""

    # Should have features
    assert d.get("ITEM_FEATURES_1") != ""

    # Provenance should exist for enriched fields
    assert len(result.provenance) >= 5
